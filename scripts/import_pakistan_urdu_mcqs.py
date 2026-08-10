#!/usr/bin/env python3
"""Import structured Pakistan Affairs and Urdu Word MCQ banks into CSS Vista."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document

from import_word_mcqs import (
    all_existing_signatures,
    clean,
    load_category,
    merge_questions,
    normalized,
    update_index,
    write_category,
)


PAKISTAN_QUESTION = re.compile(r"^(\d+)\.\s+(.+)$", re.DOTALL)
PAKISTAN_OPTION = re.compile(r"^([A-D])\.\s+(.+)$", re.DOTALL)
PAKISTAN_ANSWER = re.compile(r"^Ans\.\s*([A-D])(?:\s+Ref\..*)?$", re.IGNORECASE)
PAKISTAN_SECTION = re.compile(r"^Section\s+[IVX]+\s+-\s+(.+)$", re.IGNORECASE)

URDU_QUESTION = re.compile(r"^(\d+)\.\s+(.+)$", re.DOTALL)
URDU_OPTION = re.compile(r"^\(([A-D])\)\s+(.+)$", re.DOTALL)
URDU_ANSWER = re.compile(
    r"^درست جواب:\s*\(([A-D])\)\s*(.*?)\s*(?:\|\s*ماخذ کوڈ:.*)?$",
    re.DOTALL,
)

PAKISTAN_SLUG = "pakistan-affairs"
URDU_SLUG = "urdu-language"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def validate_sequence(questions: list[dict], expected_total: int, label: str) -> None:
    numbers = [question["source_number"] for question in questions]
    expected = list(range(1, expected_total + 1))
    if numbers != expected:
        missing = sorted(set(expected) - set(numbers))
        repeats = sorted(number for number in set(numbers) if numbers.count(number) > 1)
        raise ValueError(
            f"{label} extraction produced {len(numbers)} questions; "
            f"missing {missing[:20]}, repeated {repeats[:20]}"
        )


def extract_pakistan(path: Path) -> tuple[list[dict], dict]:
    document = Document(path)
    paragraphs = document.paragraphs
    questions: list[dict] = []
    current_topic = "Pakistan Affairs"
    index = 0

    while index < len(paragraphs):
        paragraph = paragraphs[index]
        text = paragraph.text.strip()
        style = paragraph.style.name

        if style == "Section Title Pro":
            section_match = PAKISTAN_SECTION.match(text)
            if section_match:
                current_topic = clean(section_match.group(1))
            index += 1
            continue

        if style != "MCQ Question Pro":
            index += 1
            continue

        question_match = PAKISTAN_QUESTION.match(text)
        if not question_match:
            raise ValueError(f"Unparseable Pakistan Affairs question: {text[:120]}")
        number = int(question_match.group(1))
        if index + 5 >= len(paragraphs):
            raise ValueError(f"Pakistan Affairs question {number} is incomplete")

        options: list[str] = []
        for offset, expected_letter in enumerate("ABCD", start=1):
            option_paragraph = paragraphs[index + offset]
            option_match = PAKISTAN_OPTION.match(option_paragraph.text.strip())
            if (
                option_paragraph.style.name != "MCQ Option Pro"
                or not option_match
                or option_match.group(1) != expected_letter
            ):
                raise ValueError(
                    f"Pakistan Affairs question {number} has an invalid {expected_letter} option"
                )
            options.append(clean(option_match.group(2)))

        answer_paragraph = paragraphs[index + 5]
        answer_match = PAKISTAN_ANSWER.match(answer_paragraph.text.strip())
        if answer_paragraph.style.name != "MCQ Answer Pro" or not answer_match:
            raise ValueError(f"Pakistan Affairs question {number} has an invalid answer")
        answer = ord(answer_match.group(1).upper()) - ord("A")

        question = clean(question_match.group(2))
        if not question or any(not option for option in options):
            raise ValueError(f"Pakistan Affairs question {number} contains blank content")

        questions.append(
            {
                "source_number": number,
                "q": question,
                "o": options,
                "a": answer,
                "s": current_topic,
                "d": "Intermediate",
            }
        )
        index += 6

    validate_sequence(questions, 3033, "Pakistan Affairs")
    return questions, {
        "file": path.name,
        "found": len(questions),
        "invalid": 0,
        "topics": sorted({question["s"] for question in questions}),
    }


def urdu_topics(document: Document) -> set[str]:
    topics: set[str] = set()
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if len(values) != 2:
                continue
            for value in values:
                if value and not value.isdigit() and value not in {"MCQs", "حصہ"}:
                    topics.add(clean(value))
    if not topics:
        raise ValueError("No Urdu topic headings were found")
    return topics


def extract_urdu(path: Path) -> tuple[list[dict], dict]:
    document = Document(path)
    paragraphs = document.paragraphs
    topics = urdu_topics(document)
    questions: list[dict] = []
    current_topic = "اردو زبان"
    index = 0

    while index < len(paragraphs):
        text = paragraphs[index].text.strip()
        if clean(text) in topics:
            current_topic = clean(text)
            index += 1
            continue

        question_match = URDU_QUESTION.match(text)
        if not question_match:
            index += 1
            continue

        number = int(question_match.group(1))
        if index + 5 >= len(paragraphs):
            raise ValueError(f"Urdu question {number} is incomplete")

        options: list[str] = []
        for offset, expected_letter in enumerate("ABCD", start=1):
            option_match = URDU_OPTION.match(paragraphs[index + offset].text.strip())
            if not option_match or option_match.group(1) != expected_letter:
                raise ValueError(f"Urdu question {number} has an invalid {expected_letter} option")
            options.append(clean(option_match.group(2)))

        answer_match = URDU_ANSWER.match(paragraphs[index + 5].text.strip())
        if not answer_match:
            raise ValueError(f"Urdu question {number} has an invalid answer")
        answer = ord(answer_match.group(1)) - ord("A")
        answer_text = clean(answer_match.group(2))
        if answer_text and normalized(answer_text) != normalized(options[answer]):
            raise ValueError(
                f"Urdu question {number} answer text does not match option "
                f"{answer_match.group(1)}"
            )

        question = clean(question_match.group(2))
        if not question or any(not option for option in options):
            raise ValueError(f"Urdu question {number} contains blank content")

        questions.append(
            {
                "source_number": number,
                "q": question,
                "o": options,
                "a": answer,
                "s": current_topic,
                "d": "Intermediate",
            }
        )
        index += 6

    validate_sequence(questions, 833, "Urdu")
    return questions, {
        "file": path.name,
        "found": len(questions),
        "invalid": 0,
        "topics": sorted({question["s"] for question in questions}),
    }


def reorder_categories(index: dict) -> dict:
    categories = index["categories"]
    by_slug = {category["slug"]: category for category in categories}
    ordered = [
        category
        for category in categories
        if category["slug"] not in {PAKISTAN_SLUG, URDU_SLUG}
    ]

    pakistan_at = next(
        (i + 1 for i, item in enumerate(ordered) if item["slug"] == "pakistan-history"),
        0,
    )
    ordered.insert(pakistan_at, by_slug[PAKISTAN_SLUG])

    english_at = next(
        (i + 1 for i, item in enumerate(ordered) if item["slug"] == "english-grammar"),
        len(ordered),
    )
    ordered.insert(english_at, by_slug[URDU_SLUG])
    index["categories"] = ordered
    index["total"] = sum(category["count"] for category in ordered)
    return index


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pakistan", type=Path, required=True)
    parser.add_argument("--pakistan-copy", type=Path, required=True)
    parser.add_argument("--urdu", type=Path, required=True)
    parser.add_argument("--bank-dir", type=Path, required=True)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    pakistan_hash = sha256(args.pakistan)
    pakistan_copy_hash = sha256(args.pakistan_copy)
    if pakistan_hash != pakistan_copy_hash:
        raise ValueError(
            "The two Pakistan Affairs files differ; refusing to treat either as a duplicate"
        )

    pakistan, pakistan_report = extract_pakistan(args.pakistan)
    urdu, urdu_report = extract_urdu(args.urdu)
    global_signatures = all_existing_signatures(args.bank_dir)

    existing_pakistan = load_category(args.bank_dir, PAKISTAN_SLUG)
    existing_urdu = load_category(args.bank_dir, URDU_SLUG)
    merged_pakistan, pakistan_duplicates = merge_questions(
        existing_pakistan,
        pakistan,
        PAKISTAN_SLUG,
        global_signatures,
    )
    merged_urdu, urdu_duplicates = merge_questions(
        existing_urdu,
        urdu,
        URDU_SLUG,
        global_signatures,
    )

    report = {
        "input_files": {
            "unique": [args.pakistan.name, args.urdu.name],
            "exact_duplicate_files_skipped": [args.pakistan_copy.name],
            "pakistan_sha256": pakistan_hash,
        },
        "pakistan_affairs": {
            **pakistan_report,
            "existing": len(existing_pakistan),
            "imported": len(merged_pakistan) - len(existing_pakistan),
            "duplicates_skipped": pakistan_duplicates,
        },
        "urdu_language": {
            **urdu_report,
            "existing": len(existing_urdu),
            "imported": len(merged_urdu) - len(existing_urdu),
            "duplicates_skipped": urdu_duplicates,
        },
    }

    if not args.dry_run:
        pakistan_chunks = write_category(
            args.bank_dir, PAKISTAN_SLUG, merged_pakistan
        )
        urdu_chunks = write_category(args.bank_dir, URDU_SLUG, merged_urdu)
        index = update_index(
            args.bank_dir,
            {
                PAKISTAN_SLUG: (
                    "Pakistan Affairs",
                    len(merged_pakistan),
                    pakistan_chunks,
                ),
                URDU_SLUG: (
                    "Urdu Language",
                    len(merged_urdu),
                    urdu_chunks,
                ),
            },
        )
        index = reorder_categories(index)
        (args.bank_dir / "index.json").write_text(
            json.dumps(index, ensure_ascii=False, indent=1) + "\n",
            encoding="utf-8",
        )
        report["bank_total"] = index["total"]
        report["category_count"] = len(index["categories"])

    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
