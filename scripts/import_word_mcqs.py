#!/usr/bin/env python3
"""Import structured Word MCQ banks into CSS Vista's sharded public question bank."""

from __future__ import annotations

import argparse
import json
import math
import re
from collections import defaultdict
from datetime import date
from pathlib import Path

from docx import Document


CHUNK_SIZE = 800
OPTION_LINE = re.compile(
    r"^A\.\s*(.*?)\s+B\.\s*(.*?)\s+C\.\s*(.*?)\s+D\.\s*(.*?)\s*$",
    re.DOTALL,
)
SCIENCE_ITEM = re.compile(
    r"^(\d+)\.\s*(.*?)\s+\(A\)\s*(.*?)\s+\(B\)\s*(.*?)\s+"
    r"\(C\)\s*(.*?)\s+\(D\)\s*(.*?)\s*\nAnswer:\s*([A-D])(?:\.\s*(.*))?$",
    re.DOTALL,
)
QUESTION_LINE = re.compile(r"^(\d+)\.\s*(.+)$", re.DOTALL)
ANSWER_KEY_ITEM = re.compile(r"(\d+)-([A-D])\s*\[")
SOURCE_BRANDING = re.compile(
    r"\b(?:GK\s+TABLES\s+BY\s+)?HUMZA\s+ALTAF\b",
    re.IGNORECASE,
)


def clean(value: str) -> str:
    value = SOURCE_BRANDING.sub("", value)
    return re.sub(r"\s+", " ", value).strip(" \t\r\n-|")


def normalized(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def signature(question: dict) -> str:
    return "|".join(
        [
            normalized(question["q"]),
            *(normalized(option) for option in question["o"]),
            str(question["a"]),
        ]
    )


def extract_islamiat(path: Path) -> tuple[list[dict], dict]:
    document = Document(path)
    answers: dict[int, int] = {}
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for number, letter in ANSWER_KEY_ITEM.findall(cell.text):
                    answers[int(number)] = ord(letter) - ord("A")

    questions: list[dict] = []
    paragraphs = document.paragraphs
    current_topic = "Islamic Studies"
    ignored_headings = {
        "Verification and Exclusion Standard",
        "Final Topic Distribution",
        "Answer Key with Source References",
        "Source Guide",
    }

    index = 0
    while index < len(paragraphs):
        paragraph = paragraphs[index]
        text = paragraph.text.strip()
        if paragraph.style.name == "Heading 1" and text not in ignored_headings:
            current_topic = clean(text)
            index += 1
            continue

        match = QUESTION_LINE.match(text)
        if not match:
            index += 1
            continue

        number = int(match.group(1))
        if index + 1 >= len(paragraphs):
            raise ValueError(f"Islamiat question {number} has no option paragraph")
        option_match = OPTION_LINE.match(paragraphs[index + 1].text.strip())
        if not option_match:
            raise ValueError(f"Islamiat question {number} does not contain four parseable options")
        if number not in answers:
            raise ValueError(f"Islamiat question {number} is missing from the answer key")

        question = clean(match.group(2))
        options = [clean(option) for option in option_match.groups()]
        answer = answers[number]
        if not question or len(options) != 4 or any(not option for option in options):
            raise ValueError(f"Islamiat question {number} is incomplete")

        questions.append(
            {
                "source_number": number,
                "q": question,
                "o": options,
                "a": answer,
                "s": current_topic,
                "d": "Intermediate",
            }
        )
        index += 2

    numbers = [question["source_number"] for question in questions]
    expected = list(range(1, 3221))
    if numbers != expected:
        missing = sorted(set(expected) - set(numbers))
        raise ValueError(
            f"Islamiat extraction produced {len(numbers)} questions; missing {missing[:20]}"
        )
    if sorted(answers) != expected:
        raise ValueError("Islamiat answer key is not complete from 1 through 3220")

    return questions, {
        "file": path.name,
        "found": len(questions),
        "invalid": 0,
        "topics": sorted({question["s"] for question in questions}),
    }


def extract_science(path: Path) -> tuple[list[dict], dict]:
    document = Document(path)
    questions: list[dict] = []
    current_topic = "Everyday Science"
    ignored_headings = {
        "How to Use This Bank",
        "Section Distribution",
        "Verification Framework and Authoritative References",
        "Quality-Control Summary",
    }

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if paragraph.style.name == "Heading 1":
            if text not in ignored_headings:
                current_topic = clean(re.sub(r"^Section\s+\d+:\s*", "", text))
            continue

        match = SCIENCE_ITEM.match(text)
        if not match:
            continue

        number = int(match.group(1))
        question = clean(match.group(2))
        options = [clean(match.group(i)) for i in range(3, 7)]
        answer = ord(match.group(7)) - ord("A")
        answer_text = clean(match.group(8) or "")

        if not question or len(options) != 4 or any(not option for option in options):
            raise ValueError(f"Everyday Science question {number} is incomplete")
        if answer_text and normalized(answer_text) != normalized(options[answer]):
            raise ValueError(
                f"Everyday Science question {number} answer text does not match option {match.group(7)}"
            )

        questions.append(
            {
                "source_number": number,
                "q": question,
                "o": options,
                "a": answer,
                "s": current_topic,
                "d": "Intermediate",
            }
        )

    numbers = [question["source_number"] for question in questions]
    expected = list(range(1, 3834))
    if numbers != expected:
        missing = sorted(set(expected) - set(numbers))
        raise ValueError(
            f"Everyday Science extraction produced {len(numbers)} questions; missing {missing[:20]}"
        )

    return questions, {
        "file": path.name,
        "found": len(questions),
        "invalid": 0,
        "topics": sorted({question["s"] for question in questions}),
    }


def load_category(bank_dir: Path, slug: str) -> list[dict]:
    questions: list[dict] = []
    for path in sorted(
        bank_dir.glob(f"cat-{slug}-*.json"),
        key=lambda item: int(item.stem.rsplit("-", 1)[1]),
    ):
        questions.extend(json.loads(path.read_text(encoding="utf-8")))
    return questions


def all_existing_signatures(bank_dir: Path) -> set[str]:
    signatures: set[str] = set()
    for path in bank_dir.glob("cat-*.json"):
        for question in json.loads(path.read_text(encoding="utf-8")):
            if (
                question.get("q")
                and isinstance(question.get("o"), list)
                and len(question["o"]) == 4
                and isinstance(question.get("a"), int)
            ):
                signatures.add(signature(question))
    return signatures


def max_sequence(questions: list[dict], slug: str) -> int:
    maximum = 0
    matcher = re.compile(rf"^{re.escape(slug)}-(\d+)$")
    for question in questions:
        match = matcher.match(question.get("id", ""))
        if match:
            maximum = max(maximum, int(match.group(1)))
    return maximum


def merge_questions(
    existing: list[dict],
    incoming: list[dict],
    slug: str,
    global_signatures: set[str],
) -> tuple[list[dict], int]:
    merged = list(existing)
    next_sequence = max_sequence(existing, slug) + 1
    skipped = 0

    for raw in incoming:
        question = {key: value for key, value in raw.items() if key != "source_number"}
        key = signature(question)
        if key in global_signatures:
            skipped += 1
            continue
        question["id"] = f"{slug}-{next_sequence}"
        next_sequence += 1
        merged.append(
            {
                "id": question["id"],
                "q": question["q"],
                "o": question["o"],
                "a": question["a"],
                "s": question["s"],
                "d": question["d"],
            }
        )
        global_signatures.add(key)

    return merged, skipped


def write_category(bank_dir: Path, slug: str, questions: list[dict]) -> int:
    grouped: dict[int, list[dict]] = defaultdict(list)
    for question in questions:
        sequence = int(question["id"].rsplit("-", 1)[1])
        grouped[(sequence - 1) // CHUNK_SIZE].append(question)

    for path in bank_dir.glob(f"cat-{slug}-*.json"):
        path.unlink()
    for chunk, items in sorted(grouped.items()):
        path = bank_dir / f"cat-{slug}-{chunk}.json"
        path.write_text(
            json.dumps(items, ensure_ascii=False, separators=(",", ":")),
            encoding="utf-8",
        )
    return max(grouped, default=-1) + 1


def update_index(
    bank_dir: Path,
    category_updates: dict[str, tuple[str, int, int]],
) -> dict:
    index_path = bank_dir / "index.json"
    index = json.loads(index_path.read_text(encoding="utf-8"))
    categories = index["categories"]

    for slug, (name, count, chunks) in category_updates.items():
        category = next((item for item in categories if item["slug"] == slug), None)
        if category is None:
            category = {
                "slug": slug,
                "name": name,
                "count": count,
                "chunks": chunks,
                "mpt": True,
            }
            science_index = next(
                (i for i, item in enumerate(categories) if item["slug"] == "science"),
                len(categories) - 1,
            )
            categories.insert(science_index + 1, category)
        else:
            category.update(name=name, count=count, chunks=chunks, mpt=True)

    index["generatedAt"] = date.today().isoformat()
    index["total"] = sum(category["count"] for category in categories)
    index_path.write_text(
        json.dumps(index, ensure_ascii=False, indent=1) + "\n",
        encoding="utf-8",
    )
    return index


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--islamiat", type=Path, required=True)
    parser.add_argument("--science", type=Path, required=True)
    parser.add_argument("--bank-dir", type=Path, required=True)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    islamiat, islamiat_report = extract_islamiat(args.islamiat)
    science, science_report = extract_science(args.science)
    global_signatures = all_existing_signatures(args.bank_dir)

    existing_islamiat = load_category(args.bank_dir, "islamic-gk")
    existing_science = load_category(args.bank_dir, "everyday-science")

    merged_islamiat, islamiat_duplicates = merge_questions(
        existing_islamiat, islamiat, "islamic-gk", global_signatures
    )
    merged_science, science_duplicates = merge_questions(
        existing_science, science, "everyday-science", global_signatures
    )

    report = {
        "islamiat": {
            **islamiat_report,
            "existing": len(existing_islamiat),
            "imported": len(merged_islamiat) - len(existing_islamiat),
            "duplicates_skipped": islamiat_duplicates,
        },
        "everyday_science": {
            **science_report,
            "existing": len(existing_science),
            "imported": len(merged_science) - len(existing_science),
            "duplicates_skipped": science_duplicates,
        },
    }

    if not args.dry_run:
        islamiat_chunks = write_category(
            args.bank_dir, "islamic-gk", merged_islamiat
        )
        science_chunks = write_category(
            args.bank_dir, "everyday-science", merged_science
        )
        index = update_index(
            args.bank_dir,
            {
                "islamic-gk": (
                    "Islamic General Knowledge",
                    len(merged_islamiat),
                    islamiat_chunks,
                ),
                "everyday-science": (
                    "Everyday Science",
                    len(merged_science),
                    science_chunks,
                ),
            },
        )
        report["bank_total"] = index["total"]
        report["category_count"] = len(index["categories"])

    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
